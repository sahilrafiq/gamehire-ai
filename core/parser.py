"""
core/parser.py
Read resumes from PDF/DOCX/TXT, and render AI-improved resumes back out to
BOTH .docx and .pdf -- built from the same shared, parsed structure so the
two formats always match each other. Pure Python (python-docx + reportlab),
no LibreOffice/MS Word install required for the PDF export.
"""

from dataclasses import dataclass
from pathlib import Path
from xml.sax.saxutils import escape

import pdfplumber
from docx import Document


def extract_text(path: str) -> str:
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return _from_pdf(path)
    elif ext == ".docx":
        return _from_docx(path)
    elif ext == ".txt":
        return Path(path).read_text(encoding="utf-8", errors="ignore")
    else:
        raise ValueError(f"Unsupported file type: {ext}. Use .pdf, .docx, or .txt")


def _from_pdf(path: str) -> str:
    text_parts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text_parts.append(page.extract_text() or "")
    text = "\n".join(text_parts).strip()
    if not text:
        raise ValueError(
            "No selectable text found in this PDF. It may be a scanned image — "
            "export your resume as a text-based PDF or DOCX instead."
        )
    return text


def _from_docx(path: str) -> str:
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs)


# --------------------------------------------------------------------------
# Shared structure parsing. Both the .docx and .pdf renderers below consume
# the SAME parsed blocks, so the two output formats always match each other
# and stay faithful to the AI's (structure-preserving) rewritten text.
# --------------------------------------------------------------------------

KNOWN_SECTION_NAMES = {
    "summary", "profile", "objective", "experience", "work experience",
    "professional experience", "employment", "work history", "education",
    "skills", "technical skills", "core skills", "projects", "certifications",
    "achievements", "awards",
}


@dataclass
class Block:
    kind: str  # "heading" | "bullet" | "paragraph"
    text: str


def _is_heading(line: str) -> bool:
    """Local models don't always follow the 'ALL CAPS header' instruction as
    reliably as Claude does, so heading detection also matches known resume
    section names case-insensitively, not just ALL-CAPS lines."""
    stripped = line.strip(":").strip()
    if len(stripped) > 40 or not stripped:
        return False
    if stripped.isupper():
        return True
    return stripped.lower() in KNOWN_SECTION_NAMES


def parse_blocks(text: str) -> list[Block]:
    blocks = []
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if _is_heading(line):
            blocks.append(Block("heading", line.strip(":").strip()))
        elif line.startswith(("- ", "• ", "* ")):
            blocks.append(Block("bullet", line[2:].strip()))
        else:
            blocks.append(Block("paragraph", line))
    return blocks


def text_to_docx(text: str, output_path: str) -> str:
    """Render parsed blocks into a clean, single-column, ATS-friendly .docx.
    Single-column, no tables/text-boxes, standard fonts -- the format ATS
    parsers handle best."""
    blocks = parse_blocks(text)
    doc = Document()

    for b in blocks:
        if b.kind == "heading":
            doc.add_heading(b.text, level=2)
        elif b.kind == "bullet":
            doc.add_paragraph(b.text, style="List Bullet")
        else:
            doc.add_paragraph(b.text)

    doc.save(output_path)
    return output_path


def text_to_pdf(text: str, output_path: str) -> str:
    """Render the same parsed blocks into a matching .pdf, via reportlab --
    no LibreOffice or MS Word install needed, works identically on any OS."""
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph

    blocks = parse_blocks(text)
    styles = getSampleStyleSheet()
    heading_style = ParagraphStyle(
        "ResumeHeading", parent=styles["Heading2"],
        spaceBefore=12, spaceAfter=4, textColor="#1a1a1a",
    )
    body_style = ParagraphStyle(
        "ResumeBody", parent=styles["Normal"], fontSize=10.5, leading=14, spaceAfter=4,
    )
    # Plain "- " text prefix rather than reportlab's built-in bullet glyph:
    # symbol-font bullets extract as garbage (e.g. "(cid:127)") in ATS text
    # parsers, which defeats the point of an ATS-friendly export.
    bullet_style = ParagraphStyle(
        "ResumeBullet", parent=body_style, leftIndent=14, firstLineIndent=-14,
    )

    doc = SimpleDocTemplate(
        output_path, pagesize=LETTER,
        topMargin=0.6 * inch, bottomMargin=0.6 * inch,
        leftMargin=0.75 * inch, rightMargin=0.75 * inch,
    )
    story = []

    for b in blocks:
        if b.kind == "heading":
            story.append(Paragraph(escape(b.text.upper()), heading_style))
        elif b.kind == "bullet":
            story.append(Paragraph(f"-  {escape(b.text)}", bullet_style))
        else:
            story.append(Paragraph(escape(b.text), body_style))

    doc.build(story)
    return output_path
